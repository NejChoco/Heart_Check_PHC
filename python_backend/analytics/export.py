import pandas as pd
import pytz
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

MANILA_TZ = pytz.timezone("Asia/Manila")

# Column order matches PHC's own sheet layout, so exported files can be
# handed back to PHC MIS without translation.
HEADERS = [
    "ID",
    "Queue Code",
    "Service",
    "Queuing Time",
    "Initial Assessment",
    "Doctor Seen",
    "Doctor Completed",
    "Carry Out Completed",
    "Status",
]


def _fmt_ts(ts):
    """Format a timestamp the way PHC's sheets do: MM/DD/YYYY HH:MM AM/PM, Manila time."""
    if ts is None:
        return ""
    if isinstance(ts, str):
        try:
            ts = pd.to_datetime(ts)
        except Exception:
            return ts
    if pd.isna(ts):
        return ""
    if ts.tzinfo is None:
        ts = ts.tz_localize("UTC")
    return ts.astimezone(MANILA_TZ).strftime("%m/%d/%Y %I:%M %p")


def build_phc_excel(df: pd.DataFrame, sheet_title: str = "Patients") -> BytesIO:
    """
    Builds an .xlsx matching PHC's Tracking Sheet format.

    Expects RAW patients rows (pre-preprocessing.py rename) — i.e. still has
    `service`, not `purpose`. Pull straight from Supabase, not through the
    analytics pipeline, or the column names won't line up.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]  # Excel's 31-char sheet name limit

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    row_idx = 2
    for _, row in df.iterrows():
        values = [
            row.get("id", ""),
            row.get("patientNum", "") or "",
            row.get("service", "") or "",
            _fmt_ts(row.get("created_at")),
            _fmt_ts(row.get("reg_end")),
            _fmt_ts(row.get("consult_start")),
            _fmt_ts(row.get("consult_end")),
            _fmt_ts(row.get("carryout_end")),
            row.get("status", "") or "",
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = border
            cell.alignment = Alignment(horizontal="center")
        row_idx += 1

    for col_idx, header in enumerate(HEADERS, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = max(
            [len(str(header))]
            + [len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(2, row_idx)]
        )
        ws.column_dimensions[col_letter].width = max_len + 4

    ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer